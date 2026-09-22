"""
Lee el Excel de trazabilidad y genera `datos_mes_actual.json` con la pestaña
del MES ACTUAL del sistema.

Puntos importantes de robustez (no simplificar sin leer esto):

- El Excel NO tiene nombre fijo: Brenda manda una copia nueva cada tanto y el
  nombre cambia ("Copia de Trazabilidad ... Hasta oct.xlsx"). Se busca por
  patrón en la carpeta del proyecto y se usa el archivo .xlsx de trazabilidad
  más reciente por fecha de modificación.
- El archivo vive en OneDrive y puede estar bloqueado o abierto en Excel; se
  copia a una carpeta temporal antes de leerlo para evitar PermissionError.
- El mes/año de cada pestaña se toma de la celda B2 (fecha real del mes) y solo
  si no sirve se cae al nombre de la pestaña. Así funcionan pestañas con nombre
  raro (" Enero 2026", "Septiembre" sin año, "Julio 2026 (2)").
- El estado de cada día está en el COLOR DE RELLENO de la celda, no en texto.
"""

import colorsys
import json
import re
import shutil
import sys
import tempfile
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl

BASE = Path(__file__).parent.parent
OUT_PATH = Path(__file__).parent / "datos_mes_actual.json"

MESES_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
            "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

MESES_ABREV = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
               "jul": 7, "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12}

# Orden de índice de tema usado por openpyxl en fgColor.theme
THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]

# Pestañas antiguas sin año en el nombre ni fecha en B2.
SHEET_YEAR_OVERRIDE = {"septiembre": 2025}


def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")


# --------------------------------------------------------------------------
# Ubicar y abrir el Excel
# --------------------------------------------------------------------------

def find_excel(base=BASE):
    """Devuelve el .xlsx de trazabilidad más reciente de la carpeta del proyecto."""
    candidatos = [
        p for p in base.glob("*.xlsx")
        if not p.name.startswith("~$") and "trazabilidad" in strip_accents(p.name).lower()
    ]
    if not candidatos:
        raise FileNotFoundError(
            f"No se encontró ningún archivo *.xlsx con 'Trazabilidad' en el nombre dentro de:\n  {base}\n"
            "Copia ahí el Excel que mandó Brenda (el nombre puede ser cualquiera, "
            "mientras diga 'Trazabilidad')."
        )
    return max(candidatos, key=lambda p: p.stat().st_mtime)


def copia_local(path):
    """Copia el Excel a una carpeta temporal (OneDrive/Excel suelen dejarlo bloqueado)."""
    tmpdir = Path(tempfile.mkdtemp(prefix="trazabilidad_"))
    destino = tmpdir / "libro.xlsx"
    shutil.copy2(path, destino)
    return destino


# --------------------------------------------------------------------------
# Detección de mes/año de cada pestaña
# --------------------------------------------------------------------------

def mes_desde_texto(texto):
    """'Septiembre 2026', ' Enero 2026', 'dic-2025' -> (mes, anio|None)."""
    norm = strip_accents(texto).lower().strip()
    mes = None
    palabra = re.split(r"[\s\-_]+", norm)[0] if norm else ""
    for i, m in enumerate(MESES_ES):
        if palabra == strip_accents(m).lower():
            mes = i + 1
            break
    if mes is None:
        mes = MESES_ABREV.get(palabra[:3]) if palabra[:3] in MESES_ABREV else None
    m = re.search(r"(20\d{2})", norm)
    anio = int(m.group(1)) if m else None
    return mes, anio


def fila_fechas(ws, max_scan=15):
    """Fila donde la columna E dice 'Fechas'. None si la pestaña no es de un mes."""
    for r in range(1, max_scan):
        v = ws.cell(row=r, column=5).value
        if v and strip_accents(v).strip().lower().startswith("fecha"):
            return r
    return None


def mes_anio_de_hoja(ws, nombre):
    """Mes/año de la pestaña: primero la fecha real de B2, si no el nombre."""
    b2 = ws.cell(row=2, column=2).value
    if isinstance(b2, datetime):
        return b2.month, b2.year

    mes, anio = (None, None)
    if isinstance(b2, str):
        mes, anio = mes_desde_texto(b2)
    if mes is None or anio is None:
        mes_n, anio_n = mes_desde_texto(nombre)
        mes = mes if mes is not None else mes_n
        anio = anio if anio is not None else anio_n
    if anio is None:
        anio = SHEET_YEAR_OVERRIDE.get(strip_accents(nombre).strip().lower().split()[0]
                                       if nombre.strip() else "")
    return mes, anio


def hojas_de_mes(wb):
    """[(anio, mes, orden_en_libro, nombre)] de las pestañas que sí son de un mes."""
    hojas = []
    for idx, nombre in enumerate(wb.sheetnames):
        ws = wb[nombre]
        if fila_fechas(ws) is None:
            continue  # INVENTARIO, ACTIVO FIJO, etc.
        mes, anio = mes_anio_de_hoja(ws, nombre)
        if mes is None or anio is None:
            continue
        hojas.append((anio, mes, idx, nombre))
    return hojas


def find_current_sheet(wb, hoy=None):
    """
    Elige la pestaña del mes actual del sistema.

    Si no existe (el Excel todavía no trae el mes nuevo), cae a la pestaña de mes
    más reciente que no sea futura y marca `es_mes_actual=False` para que el panel
    avise en vez de aparentar estar al día.
    """
    hoy = hoy or datetime.now()
    hojas = hojas_de_mes(wb)
    if not hojas:
        raise ValueError("El libro no tiene ninguna pestaña con estructura de mes (fila 'Fechas' en columna E).")

    exactas = [h for h in hojas if (h[0], h[1]) == (hoy.year, hoy.month)]
    if exactas:
        # Puede haber duplicados tipo "Julio 2026 (2)": se prefiere el nombre sin
        # sufijo y, entre iguales, la pestaña que aparece más a la derecha.
        exactas.sort(key=lambda h: (bool(re.search(r"\(\s*\d+\s*\)", h[3])), -h[2]))
        anio, mes, _, nombre = exactas[0]
        return nombre, MESES_ES[mes - 1], anio, True

    pasadas = [h for h in hojas if (h[0], h[1]) <= (hoy.year, hoy.month)]
    candidatas = pasadas or hojas
    anio, mes, _, nombre = max(candidatas, key=lambda h: (h[0], h[1], h[2]))
    return nombre, MESES_ES[mes - 1], anio, False


# --------------------------------------------------------------------------
# Color -> estado
# --------------------------------------------------------------------------

def get_theme_palette(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/theme/theme1.xml").decode("utf-8")
    m = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S)
    scheme = m.group(0)
    found = re.findall(
        r'<a:(\w+)>.*?(?:srgbClr val="([0-9A-Fa-f]{6})"|sysClr[^>]*lastClr="([0-9A-Fa-f]{6})")', scheme
    )
    colors = {name: (rgb1 or rgb2) for name, rgb1, rgb2 in found}
    return [colors[name] for name in THEME_ORDER]


def apply_tint(rgb_hex, tint):
    r = int(rgb_hex[0:2], 16) / 255
    g = int(rgb_hex[2:4], 16) / 255
    b = int(rgb_hex[4:6], 16) / 255
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    if tint < 0:
        l = l * (1.0 + tint)
    else:
        l = l * (1.0 - tint) + tint
    r2, g2, b2 = colorsys.hls_to_rgb(h, l, s)
    return (round(r2 * 255), round(g2 * 255), round(b2 * 255))


def color_distance(c1, c2):
    return sum(abs(a - b) for a, b in zip(c1, c2))


class ColorClassifier:
    # Mapeo verificado empíricamente (Claude_Prompt.md) + confirmado matemáticamente
    # contra el theme real del archivo: theme=3/tint=0.75, theme=9/tint=0.6, theme=5/tint=0.6.
    # Algunas celdas quedan con el color "horneado" a RGB explícito en vez de theme+tint
    # (pasa al editar en Excel) -- por eso comparamos por distancia de color, no por tipo exacto.
    UMBRAL_DISTANCIA = 15

    def __init__(self, theme_palette):
        self.theme_palette = theme_palette
        self.refs = {
            "Demostración": apply_tint(theme_palette[3], 0.75),
            "Préstamo": apply_tint(theme_palette[9], 0.6),
            "Incompleto": apply_tint(theme_palette[5], 0.6),
        }

    def classify(self, cell):
        fill = cell.fill
        if not fill or fill.patternType is None:
            return None
        fg = fill.fgColor
        if fg.type == "theme":
            rgb = apply_tint(self.theme_palette[fg.theme], fg.tint or 0.0)
        elif fg.type == "rgb" and fg.rgb and len(str(fg.rgb)) >= 6:
            hexval = str(fg.rgb)[-6:]
            rgb = (int(hexval[0:2], 16), int(hexval[2:4], 16), int(hexval[4:6], 16))
        else:
            return None
        mejor, mejor_dist = None, 9999
        for estado, ref in self.refs.items():
            d = color_distance(rgb, ref)
            if d < mejor_dist:
                mejor, mejor_dist = estado, d
        if mejor_dist <= self.UMBRAL_DISTANCIA:
            return mejor
        return "DESCONOCIDO"


# --------------------------------------------------------------------------
# Parseo de la pestaña del mes
# --------------------------------------------------------------------------

def limpio(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def parse_month_sheet(ws, classifier):
    fechas_row = fila_fechas(ws)
    if fechas_row is None:
        raise ValueError("No se encontró la fila 'Fechas' en la columna E de esta pestaña")

    day_number_row = fechas_row + 2
    data_start_row = fechas_row + 3

    day_cols = []
    c = 5
    while True:
        v = ws.cell(row=day_number_row, column=c).value
        if isinstance(v, (int, float)) and int(v) == len(day_cols) + 1:
            day_cols.append(c)
            c += 1
        else:
            break
    num_dias = len(day_cols)
    if num_dias == 0:
        raise ValueError("No se detectaron columnas de días bajo la fila 'Fechas'")

    buckets = {"Demostración": [], "Préstamo": [], "Incompleto": [], "Nota/Servicio": [], "Disponible": []}
    avisos = []

    categoria_actual = None
    r = data_start_row
    while r <= ws.max_row:
        a = limpio(ws.cell(row=r, column=1).value)
        b = limpio(ws.cell(row=r, column=2).value)
        c_ = limpio(ws.cell(row=r, column=3).value)
        d = limpio(ws.cell(row=r, column=4).value)

        if a and not b and not c_ and not d:
            categoria_actual = a
            r += 1
            continue

        if b:
            modelo, serie, descripcion = b, (c_ or ""), (d or "")
            day_entries = []
            for i, col in enumerate(day_cols):
                dia = i + 1
                cell = ws.cell(row=r, column=col)
                estado = classifier.classify(cell)
                texto = limpio(cell.value)
                if estado in ("Demostración", "Préstamo", "Incompleto"):
                    day_entries.append((dia, estado, texto or ""))
                elif estado == "DESCONOCIDO":
                    avisos.append(f"Fila {r} (equipo {modelo}), día {dia}: color de relleno no reconocido")
                    day_entries.append((dia, None, None))
                elif texto:
                    day_entries.append((dia, "Nota/Servicio", texto))
                else:
                    day_entries.append((dia, None, None))

            segmentos_por_estado = {}
            i = 0
            while i < len(day_entries):
                dia, tipo, detalle = day_entries[i]
                if tipo is None:
                    i += 1
                    continue
                inicio = fin = dia
                j = i + 1
                while j < len(day_entries) and day_entries[j][1] == tipo and day_entries[j][2] == detalle:
                    fin = day_entries[j][0]
                    j += 1
                segmentos_por_estado.setdefault(tipo, []).append({"lugar": detalle, "inicio": inicio, "fin": fin})
                i = j

            registro_base = {"categoria": categoria_actual, "modelo": modelo, "numero_serie": serie, "descripcion": descripcion}
            if not segmentos_por_estado:
                buckets["Disponible"].append({**registro_base, "segmentos": []})
            else:
                for tipo, segs in segmentos_por_estado.items():
                    buckets[tipo].append({**registro_base, "segmentos": segs})
        r += 1

    return num_dias, buckets, avisos


def main():
    excel_original = find_excel()
    excel = copia_local(excel_original)

    wb = openpyxl.load_workbook(excel, data_only=True)
    sheet_name, mes_nombre, anio, es_mes_actual = find_current_sheet(wb)
    ws = wb[sheet_name]
    classifier = ColorClassifier(get_theme_palette(excel))

    num_dias, buckets, avisos = parse_month_sheet(ws, classifier)

    data = {
        "mes": mes_nombre,
        "anio": anio,
        "num_dias": num_dias,
        "es_mes_actual": es_mes_actual,
        "hoja": sheet_name,
        "archivo_origen": excel_original.name,
        "generado_en": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "buckets": buckets,
    }

    # La tarea programada corre todos los días, pero el Excel cambia pocas veces.
    # Si los datos son idénticos a los de la corrida anterior se conserva la marca
    # de tiempo vieja: así el archivo queda byte a byte igual, git no ve cambios y
    # no se acumulan commits vacíos. De paso, la fecha que muestra el panel pasa a
    # significar "datos actualizados al", que es lo que realmente importa.
    if OUT_PATH.exists():
        try:
            previo = json.loads(OUT_PATH.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            previo = None
        if previo and {k: v for k, v in previo.items() if k != "generado_en"} == \
                      {k: v for k, v in data.items() if k != "generado_en"}:
            data["generado_en"] = previo.get("generado_en", data["generado_en"])

    OUT_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"Excel usado   : {excel_original.name}")
    print(f"Pestaña usada : {sheet_name}")
    print(f"Mes/año       : {mes_nombre} {anio}  |  días en el mes: {num_dias}")
    if not es_mes_actual:
        print(
            f"AVISO: el Excel todavía no tiene pestaña del mes actual "
            f"({MESES_ES[datetime.now().month - 1]} {datetime.now().year}); "
            f"se está usando la última pestaña disponible ({sheet_name}) como respaldo."
        )
    for estado, items in buckets.items():
        print(f"  {estado}: {len(items)} equipos")
    if avisos:
        print(f"\n{len(avisos)} AVISO(S) - colores no reconocidos, revisar manualmente:")
        for av in avisos[:20]:
            print("  -", av)
    print(f"\nGuardado en: {OUT_PATH}")

    shutil.rmtree(excel.parent, ignore_errors=True)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        raise SystemExit(1)
